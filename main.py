import os
import sys
import json
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from table_handler import create_table, calculate_font_size, fill_table, fill_table_reversed_cols
from gen_primerov import get_examples

def get_base_dir():

    if getattr(sys, 'frozen', False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))

def load_config():

    base_dir = get_base_dir()
    config_path = os.path.join(base_dir, "config.json")
    
    if not os.path.exists(config_path):
        default_config = {
            "drobi": [[5, "+"], [5, "-"]],
            "power": [[5, "^"], [5, "root"]],
            "linear": [[5, "+"], [5, "-"]]
        }
        with open(config_path, "w", encoding="utf-8") as f:
            json.dump(default_config, f, ensure_ascii=False, indent=2)
        print(f"Создан дефолтный config.json: {config_path}")
        return default_config
    
    with open(config_path, "r", encoding="utf-8") as f:
        return json.load(f)

def create_page_pair(doc, examples_batch, answers_batch, rows, cols, 
                     cell_width_cm, cell_height_cm, page_num):

    font_size = calculate_font_size(cell_width_cm, cell_height_cm)
    
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(f"Примеры - Лист {page_num}")
    run.font.size = Pt(14)
    run.font.bold = True
    title.paragraph_format.space_after = Pt(2)
    title.paragraph_format.space_before = Pt(0)
    doc.add_paragraph()
    
    table1 = create_table(doc, rows, cols, cell_width_cm, cell_height_cm)
    fill_table(table1, examples_batch, font_size, bold=False)
    
    doc.add_page_break()
    
    title2 = doc.add_paragraph()
    title2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = title2.add_run(f"Ответы - Лист {page_num}")
    run2.font.size = Pt(14)
    run2.font.bold = True
    title.paragraph_format.space_after = Pt(2)
    title.paragraph_format.space_before = Pt(0)
    doc.add_paragraph()
    
    table2 = create_table(doc, rows, cols, cell_width_cm, cell_height_cm)
    fill_table_reversed_cols(table2, answers_batch, font_size, bold=True)
    
    #doc.add_page_break()

def create_math_document(config, output_file,
                         rows=8, cols=4,
                         cell_width_cm=4.0, cell_height_cm=3.0):

    all_examples = get_examples(config)
    examples_list = list(all_examples.keys())
    answers_list = list(all_examples.values())
    
    total_examples = len(examples_list)
    cells_per_page = rows * cols
    
    doc = Document()
    
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    
    page_num = 1
    for start_idx in range(0, total_examples, cells_per_page):
        end_idx = min(start_idx + cells_per_page, total_examples)
        
        batch_examples = examples_list[start_idx:end_idx]
        batch_answers = answers_list[start_idx:end_idx]
        
        while len(batch_examples) < cells_per_page:
            batch_examples.append("")
            batch_answers.append("")
        
        create_page_pair(
            doc, batch_examples, batch_answers,
            rows, cols, cell_width_cm, cell_height_cm, page_num
        )
        page_num += 1
    
    doc.save(output_file)
    return output_file

if __name__ == "__main__":
    config = load_config()
    print(f"Загружен конфиг: {list(config.keys())}")
    
    rows_input = input("Строки [8]: ")
    rows = int(rows_input) if rows_input.strip() else 8
    
    cols_input = input("Колонки [4]: ")
    cols = int(cols_input) if cols_input.strip() else 4
    
    width_input = input("Ширина ячейки (см) [4.0]: ")
    width = float(width_input) if width_input.strip() else 4.0
    
    height_input = input("Высота ячейки (см) [3.0]: ")
    height = float(height_input) if height_input.strip() else 3.0
    
    base_dir = get_base_dir()
    default_name = "примеры.docx"
    filename_input = input(f"Имя файла [{default_name}]: ").strip()
    filename = filename_input if filename_input else default_name
    output_path = os.path.join(base_dir, filename)
    
    create_math_document(
        config=config,
        output_file=output_path,
        rows=rows,
        cols=cols,
        cell_width_cm=width,
        cell_height_cm=height
    )