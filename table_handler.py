from docx import Document
from docx.shared import Cm, Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_margins(cell, top=0, left=0, bottom=0, right=0):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    
    for name, val in [('top', top), ('left', left), ('bottom', bottom), ('right', right)]:
        margin = OxmlElement(f'w:{name}')
        margin.set(qn('w:w'), str(val))
        margin.set(qn('w:type'), 'dxa')
        tcMar.append(margin)
    
    tcPr.append(tcMar)


def create_table(doc, rows, cols, cell_width_cm, cell_height_cm):
    table = doc.add_table(rows=rows, cols=cols)
    table.style = 'Table Grid'
    table.autofit = False
    table.allow_autofit = False
    
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement('w:tblPr')
    jc = OxmlElement('w:jc')
    jc.set(qn('w:val'), 'center')
    tblPr.append(jc)

    tblGrid = tbl.tblGrid
    
    for gridCol in tblGrid.gridCol_lst:
        tblGrid.remove(gridCol)
    
    width_twips = int(cell_width_cm * 567)
    for _ in range(cols):
        gridCol = OxmlElement('w:gridCol')
        gridCol.set(qn('w:w'), str(width_twips))
        tblGrid.append(gridCol)
    
    for row in table.rows:
        tr = row._tr
        trPr = tr.get_or_add_trPr()
        trHeight = OxmlElement('w:trHeight')
        trHeight.set(qn('w:val'), str(int(cell_height_cm * 567)))
        trHeight.set(qn('w:hRule'), "exact")
        trPr.append(trHeight)
        
        for cell in row.cells:

            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            tcW = OxmlElement('w:tcW')
            tcW.set(qn('w:w'), str(width_twips))
            tcW.set(qn('w:type'), 'dxa')
            tcPr.append(tcW)

            set_cell_margins(cell, 50, 50, 50, 50)

    return table


def calculate_font_size(cell_width_cm, cell_height_cm):
    font_size = int(12 * (cell_width_cm / 4.0) * (cell_height_cm / 3.0))
    return Pt(max(9, min(font_size, 20)))


def set_cell(table, x, y, value, font_size, bold=False, align='center'):
    cell = table.cell(x, y)
    cell.text = ""
    
    paragraph = cell.paragraphs[0]
    
    align_map = {
        'left': WD_ALIGN_PARAGRAPH.LEFT,
        'center': WD_ALIGN_PARAGRAPH.CENTER,
        'right': WD_ALIGN_PARAGRAPH.RIGHT,
    }
    paragraph.alignment = align_map.get(align, WD_ALIGN_PARAGRAPH.CENTER)
    
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    vAlign = OxmlElement('w:vAlign')
    vAlign.set(qn('w:val'), 'center')
    tcPr.append(vAlign)

    run = paragraph.add_run(str(value))
    run.font.size = font_size
    run.font.bold = bold
    
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.line_spacing = Pt(0)


def fill_table(table, data, font_size, bold=False):
    rows = len(table.rows)
    cols = len(table.columns)
    
    for i, value in enumerate(data):
        if i >= rows * cols:
            break
        x = i // cols
        y = i % cols
        set_cell(table, x, y, value, font_size, bold=bold)


def fill_table_reversed_cols(table, data, font_size, bold=True):
    rows = len(table.rows)
    cols = len(table.columns)
    
    for i, value in enumerate(data):
        if i >= rows * cols:
            break
        
        x = i // cols
        original_y = i % cols
        
        y = (cols - 1) - original_y
        
        set_cell(table, x, y, value, font_size, bold=bold)